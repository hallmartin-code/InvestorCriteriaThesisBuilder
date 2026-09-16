"""Turn an uploaded deck into content blocks for one model call (CLAUDE.md §8).

A PDF is sent whole as a native `document` block, so the model sees charts and layout rather
than re-flattened text. PPTX and DOCX are read as text, because the Railpack image has no
LibreOffice; slide numbers come from the file itself, and DOCX cites sections instead.
"""

from __future__ import annotations

import base64
import hashlib
import io
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

SUPPORTED = (".pdf", ".pptx", ".docx")
LEGACY = (".ppt", ".doc", ".key")
MAX_PAGES = 200
DOCX_PARAGRAPHS_PER_SECTION = 25


class DeckError(ValueError):
    """The upload is not a deck this app can read."""


@dataclass
class Deck:
    filename: str
    kind: str
    page_count: int
    blocks: list[dict[str, Any]]
    content_hash: str
    warnings: list[str] = field(default_factory=list)

    @property
    def unit(self) -> str:
        return "section" if self.kind == "docx" else "slide"


def load_deck(filename: str, data: bytes) -> Deck:
    suffix = Path(filename or "").suffix.lower()
    if suffix in LEGACY:
        raise DeckError(f"{suffix} files are not supported here. Save the deck as .pptx or PDF and try again.")
    if suffix not in SUPPORTED:
        raise DeckError("Upload a .pdf, .pptx or .docx deck.")
    if not data:
        raise DeckError("That file is empty.")
    digest = hashlib.sha256(data).hexdigest()
    if suffix == ".pdf":
        return _pdf(filename, data, digest)
    if suffix == ".pptx":
        return _pptx(filename, data, digest)
    return _docx(filename, data, digest)


def _pdf(filename: str, data: bytes, digest: str) -> Deck:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise DeckError("That PDF is password-protected. Upload an unlocked copy.")
        pages = len(reader.pages)
    except DeckError:
        raise
    except (PdfReadError, OSError, ValueError) as exc:
        raise DeckError("That PDF could not be read. It may be corrupt.") from exc

    warnings: list[str] = []
    if pages > MAX_PAGES:
        raise DeckError(f"That deck has {pages} pages; the limit is {MAX_PAGES}.")
    if pages == 0:
        raise DeckError("That PDF has no pages.")
    if not _pdf_has_text(data):
        warnings.append("The PDF has little extractable text; it may be a scan. Slide citations may be approximate.")

    blocks = [
        {"type": "document",
         "source": {"type": "base64", "media_type": "application/pdf",
                    "data": base64.standard_b64encode(data).decode("ascii")}},
        {"type": "text", "text": f"The deck above is {filename} with {pages} slides, numbered 1 to {pages}."},
    ]
    return Deck(filename=filename, kind="pdf", page_count=pages, blocks=blocks, content_hash=digest, warnings=warnings)


def _pdf_has_text(data: bytes, minimum: int = 200) -> bool:
    try:
        import pymupdf

        with pymupdf.open(stream=data, filetype="pdf") as document:
            found = 0
            for page in document:
                found += len(page.get_text().strip())
                if found >= minimum:
                    return True
    except Exception:  # text detection is advisory only
        return True
    return False


def _pptx(filename: str, data: bytes, digest: str) -> Deck:
    from pptx import Presentation

    try:
        deck = Presentation(io.BytesIO(data))
    except Exception as exc:
        raise DeckError("That PowerPoint file could not be read. It may be corrupt.") from exc

    slides: list[str] = []
    for index, slide in enumerate(deck.slides, start=1):
        parts: list[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                parts.append(shape.text_frame.text.strip())
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    parts.append(" | ".join(cell.text.strip() for cell in row.cells))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            parts.append("Speaker notes: " + slide.notes_slide.notes_text_frame.text.strip())
        slides.append(f"--- SLIDE {index} ---\n" + ("\n".join(parts) if parts else "(no text on this slide)"))

    if not slides:
        raise DeckError("That PowerPoint file has no slides.")
    warnings = ["Slide images were not sent: this server cannot render PowerPoint. Charts and diagrams "
                "are only read through their text."]
    blocks = [{"type": "text", "text": f"Deck: {filename} ({len(slides)} slides)\n\n" + "\n\n".join(slides)}]
    return Deck(filename=filename, kind="pptx", page_count=len(slides), blocks=blocks,
                content_hash=digest, warnings=warnings)


def _docx(filename: str, data: bytes, digest: str) -> Deck:
    import docx

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise DeckError("That Word file could not be read. It may be corrupt.") from exc

    paragraphs = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            paragraphs.append(" | ".join(cell.text.strip() for cell in row.cells))
    if not paragraphs:
        raise DeckError("That Word file has no text.")

    sections: list[str] = []
    for index, start in enumerate(range(0, len(paragraphs), DOCX_PARAGRAPHS_PER_SECTION), start=1):
        body = "\n".join(paragraphs[start:start + DOCX_PARAGRAPHS_PER_SECTION])
        sections.append(f"--- SECTION {index} ---\n{body}")
    warnings = ["This is a document, not a slide deck: citations refer to numbered sections."]
    blocks = [{"type": "text", "text": f"Document: {filename} ({len(sections)} sections)\n\n" + "\n\n".join(sections)}]
    return Deck(filename=filename, kind="docx", page_count=len(sections), blocks=blocks,
                content_hash=digest, warnings=warnings)
