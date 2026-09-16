"""Deck ingestion and the screening pipeline (CLAUDE.md §8)."""

from __future__ import annotations

import io

import pytest

from icb.ingest.router import DeckError, load_deck
from icb.llm.client import ModelResult
from icb.screen import extract
from icb.screen.models import ScreeningExtraction
from test_criteria import make_pack


def make_pdf(pages: int = 3) -> bytes:
    import pymupdf

    document = pymupdf.open()
    for index in range(pages):
        page = document.new_page()
        page.insert_text((72, 96), f"Slide {index + 1}: revenue and team details for the round.")
    data = document.tobytes()
    document.close()
    return data


def make_pptx(slides: int = 2) -> bytes:
    from pptx import Presentation
    from pptx.util import Inches

    deck = Presentation()
    for index in range(slides):
        slide = deck.slides.add_slide(deck.slide_layouts[5])
        slide.shapes.title.text = f"Slide {index + 1}"
        box = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(6), Inches(1))
        box.text_frame.text = "ARR is $1.2M with 40 paying customers."
    buffer = io.BytesIO()
    deck.save(buffer)
    return buffer.getvalue()


def make_docx(paragraphs: int = 30) -> bytes:
    import docx

    document = docx.Document()
    for index in range(paragraphs):
        document.add_paragraph(f"Paragraph {index + 1}: the company sells to mid-market buyers.")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def extraction_payload(score: int | None = 4) -> dict:
    return ScreeningExtraction(
        factors=[{"factor_id": f"factor_{i}", "score": score, "evidence_standard_met": score is not None,
                  "evidence": [{"quote": "ARR is $1.2M", "slide": 2, "evidence_type": "deck_statement"}]}
                 for i in range(1, 7)],
        hard_criteria=[{"criterion_id": "stage", "result": "MET", "deck_evidence": "Seed round stated",
                        "classification": "FACT", "slides": [1], "excerpt": "Seed round"}],
        screening_summary="Analytical summary.", thesis_fit="Fits the stated thesis.",
    ).model_dump(mode="json")


def test_pdf_is_sent_whole_with_its_page_count():
    deck = load_deck("deck.pdf", make_pdf(4))
    assert deck.kind == "pdf" and deck.page_count == 4 and deck.unit == "slide"
    assert deck.blocks[0]["type"] == "document"
    assert deck.blocks[0]["source"]["media_type"] == "application/pdf"


def test_pptx_and_docx_are_read_as_text():
    pptx = load_deck("deck.pptx", make_pptx(3))
    assert pptx.page_count == 3 and pptx.blocks[0]["type"] == "text"
    assert "--- SLIDE 2 ---" in pptx.blocks[0]["text"] and "ARR is $1.2M" in pptx.blocks[0]["text"]
    assert pptx.warnings  # no slide images on this server

    docx_deck = load_deck("notes.docx", make_docx(30))
    assert docx_deck.unit == "section" and docx_deck.page_count == 2
    assert "--- SECTION 1 ---" in docx_deck.blocks[0]["text"]


@pytest.mark.parametrize(("filename", "data", "fragment"), [
    ("deck.ppt", b"anything", "not supported here"),
    ("deck.key", b"anything", "not supported here"),
    ("deck.exe", b"MZ", "Upload a .pdf"),
    ("deck.pdf", b"", "empty"),
    ("deck.pdf", b"not a pdf at all", "could not be read"),
])
def test_unreadable_uploads_are_refused(filename, data, fragment):
    with pytest.raises(DeckError, match=fragment):
        load_deck(filename, data)


def test_citations_outside_the_deck_are_rejected():
    deck = load_deck("deck.pdf", make_pdf(2))
    validate = extract._validator(deck)
    payload = extraction_payload()
    validate(payload)  # slide 2 exists

    payload["factors"][0]["evidence"][0]["slide"] = 9
    with pytest.raises(ValueError, match="citations cannot exist: 9"):
        validate(payload)


def test_screening_applies_the_pack_and_caches(tmp_path, monkeypatch):
    monkeypatch.setenv("ICB_CACHE_DIR", str(tmp_path / "cache"))
    pack = make_pack()
    calls = []

    def fake_call(**kwargs):
        calls.append(kwargs)
        return ModelResult(data=extraction_payload(), model="claude-opus-5",
                           input_tokens=100, output_tokens=50, corrected=False)

    monkeypatch.setattr(extract.llm, "call_json", fake_call)
    data = make_pdf(3)
    first = extract.screen_deck(slug="acme-capital", investor_name="Acme Capital",
                                filename="deck.pdf", data=data, pack=pack)

    assert first.decision.decision == "ADVANCE" and first.decision.rule == 4
    assert first.extraction.factors[0].label == "Factor 1"      # label copied from the pack
    assert first.extraction.factors[0].weight == 20
    assert first.provenance.criteria_pack["hash"] == pack.content_hash()
    assert first.provenance.cached is False

    second = extract.screen_deck(slug="acme-capital", investor_name="Acme Capital",
                                 filename="deck.pdf", data=data, pack=pack)
    assert len(calls) == 1 and second.provenance.cached is True
    assert second.decision.rule_sentence == first.decision.rule_sentence


def test_unscorable_factors_hold_the_deal_instead_of_scoring_low(tmp_path, monkeypatch):
    monkeypatch.setenv("ICB_CACHE_DIR", str(tmp_path / "cache"))
    monkeypatch.setattr(extract.llm, "call_json", lambda **kwargs: ModelResult(
        data=extraction_payload(score=None), model="claude-opus-5", input_tokens=1, output_tokens=1, corrected=False))
    screening = extract.screen_deck(slug="acme-capital", investor_name="Acme Capital",
                                    filename="deck.pdf", data=make_pdf(2), pack=make_pack())
    assert screening.decision.decision == "HOLD — REQUEST EVIDENCE"
    assert screening.decision.weighted_score is None and screening.decision.evidence_coverage == 0
